from common171b import *
import openpyxl

books=jr(O/'QA/RESULTS_WORKBOOK_PAYLOAD.json');checks=[];total=0
for b in books:
    wb=openpyxl.load_workbook(O/(b['name']+'.xlsx'),read_only=True,data_only=False)
    assert wb.sheetnames==[s['name'] for s in b['sheets']]
    for s in b['sheets']:
        expected=[s['columns']]+s['rows'];actual=list(wb[s['name']].iter_rows(values_only=True));assert len(actual)==len(expected)
        count=0
        for a,e in zip(actual,expected):
            assert len(a)==len(e)
            for x,y in zip(a,e):
                assert ('' if x is None else x)==y,(b['name'],s['name'],x,y)
                assert not(isinstance(x,str) and x in ['#REF!','#DIV/0!','#VALUE!','#NAME?','#NUM!','#SPILL!'])
                count+=1
        total+=count;checks.append(dict(workbook=b['name'],sheet=s['name'],cells=count,all_cells_match=True))
    wb.close()
csvwrite(O/'QA/RESULT_WORKBOOK_ALL_CELL_READBACK.csv',checks)
js(O/'QA/RESULT_WORKBOOK_READBACK_VERIFICATION.json',dict(utc=now(),workbooks=len(books),sheets=len(checks),cells=total,pass_all=True,visual_review='PENDING_HUMAN_OR_MODEL_IMAGE_INSPECTION'))
event('RESULT_WORKBOOK_ALL_CELL_READBACK_PASSED',workbooks=len(books),cells=total,visual_review_pending=True)
